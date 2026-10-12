"""Oops Mac 助手的固定 PPT 渲染器：outline.json → 16:9 pptx。模型只写大纲，不再自己写代码运行。

outline.json：{"title": "封面标题", "subtitle": "副标题（可空）", "slides": [{"title": "页标题", "bullets": ["要点", ...]}, ...]}
最多 12 页内容、每页 6 条要点、每条 120 字；超出部分截掉。
运行：~/.oops-agent/venv/bin/python -I oops_slides.py outline.json slides.pptx
"""
import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT = "Microsoft YaHei"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x2E, 0x86, 0xC1)
DARK = RGBColor(0x33, 0x33, 0x33)
LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xD6, 0xE4, 0xF0)
GREY = RGBColor(0x88, 0x88, 0x88)
MAX_SLIDES, MAX_BULLETS, MAX_CHARS = 12, 6, 120


def clip(value, limit):
    value = " ".join(str(value or "").split())
    return value if len(value) <= limit else value[:limit - 1] + "…"


def read_outline(path):
    if os.path.islink(path) or os.path.getsize(path) > 200_000:
        raise ValueError("大纲文件无效")
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    if not isinstance(data, dict) or not isinstance(data.get("slides"), list):
        raise ValueError("大纲缺少 slides")
    slides = []
    for item in data["slides"][:MAX_SLIDES]:
        if not isinstance(item, dict):
            continue
        bullets = [clip(b, MAX_CHARS) for b in (item.get("bullets") or [])[:MAX_BULLETS] if clip(b, MAX_CHARS)]
        if clip(item.get("title"), 30):
            slides.append(dict(title=clip(item["title"], 30), bullets=bullets))
    if not slides:
        raise ValueError("大纲里没有可用的页")
    return clip(data.get("title"), 40) or slides[0]["title"], clip(data.get("subtitle"), 80), slides


def styled(run, text, size, color, bold=False):
    run.text = text
    run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()  # East Asian glyphs take their font from <a:ea>
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", FONT)


def rectangle(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def cover(prs, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rectangle(s, 0, 0, prs.slide_width, prs.slide_height, NAVY)
    rectangle(s, Inches(0.9), Inches(3.55), Inches(1.6), Inches(0.08), ACCENT)
    box = s.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.4)).text_frame
    box.word_wrap = True
    styled(box.paragraphs[0].add_run(), title, 44, WHITE, True)
    if subtitle:
        sub = s.shapes.add_textbox(Inches(0.9), Inches(3.9), Inches(11.5), Inches(1.6)).text_frame
        sub.word_wrap = True
        styled(sub.paragraphs[0].add_run(), subtitle, 22, PALE)


def content(prs, page, index, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rectangle(s, 0, 0, prs.slide_width, prs.slide_height, LIGHT)
    rectangle(s, 0, 0, prs.slide_width, Inches(1.25), NAVY)
    head = s.shapes.add_textbox(Inches(0.7), Inches(0.2), Inches(12), Inches(0.9)).text_frame
    head.vertical_anchor = MSO_ANCHOR.MIDDLE
    styled(head.paragraphs[0].add_run(), page["title"], 32, WHITE, True)
    body = s.shapes.add_textbox(Inches(0.9), Inches(1.7), Inches(11.6), Inches(5.2)).text_frame
    body.word_wrap = True
    size = 24 if len(page["bullets"]) <= 4 else 20
    for i, text in enumerate(page["bullets"]):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.space_after = Pt(14)
        styled(p.add_run(), "■  ", size - 8, ACCENT, True)
        styled(p.add_run(), text, size, DARK)
    foot = s.shapes.add_textbox(Inches(11.8), Inches(6.95), Inches(1.2), Inches(0.4)).text_frame.paragraphs[0]
    foot.alignment = PP_ALIGN.RIGHT
    styled(foot.add_run(), f"{index} / {total}", 12, GREY)


def render(outline_path, out_path):
    title, subtitle, slides = read_outline(outline_path)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    cover(prs, title, subtitle)
    for i, page in enumerate(slides, 2):
        content(prs, page, i, len(slides) + 1)
    prs.save(out_path)
    return len(slides) + 1


if __name__ == "__main__":
    print(f"生成 {render(sys.argv[1], sys.argv[2])} 页")
